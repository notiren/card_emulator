import serial
import time
from openpyxl import load_workbook, Workbook
from filelock import FileLock
import threading
import queue
import ctypes
import os
from serial.tools import list_ports

# --- RF IDeas SDK CLASS ---
class RFIdeasReader:
    def __init__(self, dll_path="pcProxAPI.dll"):
        # Get script directory
        script_dir = os.path.dirname(os.path.abspath(__file__))
        full_path = os.path.join(script_dir, dll_path)
        
        try:
            self.dll = ctypes.windll.LoadLibrary(full_path)
            print(f"  (loaded from {full_path})")
        except OSError as e:
            raise Exception(f"Failed to load {full_path}: {e}")
        
        self.buffer_size = 256
        self.running = False

        # Find and expose functions
        self._expose_functions()

    def _expose_functions(self):
        """Find and expose DLL functions from actual exports"""
        # Need to find the correct lowercase function names
        exposed = {}
        
        # Try both uppercase and lowercase variants
        all_functions = [
            ('com_connect', ['comConnect', 'ComConnect']),
            ('com_connect_port', ['comConnectPort', 'ComConnectPort']),
            ('usb_disconnect', ['USBDisconnect', 'usbDisconnect']),
            ('com_disconnect', ['ComDisconnect', 'comDisconnect']),
            ('ping', ['Ping', 'ping']),
            ('get_dev_cnt', ['getDevCnt', 'GetDevCnt']),
            ('set_act_dev', ['setActDev', 'SetActDev']),
            ('read_cfg', ['readCfg', 'ReadCfg']),
            ('set_dev_type_srch', ['SetDevTypeSrch', 'setDevTypeSrch']),
            ('set_connect_product', ['SetConnectProduct', 'setConnectProduct']),
            ('get_product', ['GetProduct', 'getProduct']),
            ('get_id', ['getActiveID32', 'GetActiveID32', 'getActiveID', 'GetActiveID']),
            ('get_byte', ['getActiveID_byte', 'GetActiveID_byte'])
        ]
        
        for key, names in all_functions:
            for func_name in names:
                try:
                    func = getattr(self.dll, func_name)
                    
                    # Set proper signatures based on function
                    if key == 'com_connect':
                        func.argtypes = []
                        func.restype = ctypes.c_ushort
                    elif key == 'com_connect_port':
                        func.argtypes = [ctypes.c_ushort]
                        func.restype = ctypes.c_ushort
                    elif key == 'usb_disconnect':
                        func.argtypes = []
                        func.restype = ctypes.c_ushort
                    elif key == 'com_disconnect':
                        func.argtypes = []
                        func.restype = ctypes.c_ushort
                    elif key == 'ping':
                        func.argtypes = []
                        func.restype = ctypes.c_int
                    elif key == 'get_dev_cnt':
                        func.argtypes = []
                        func.restype = ctypes.c_short
                    elif key == 'set_act_dev':
                        func.argtypes = [ctypes.c_short]
                        func.restype = ctypes.c_ushort
                    elif key == 'read_cfg':
                        func.argtypes = []
                        func.restype = ctypes.c_ushort
                    elif key == 'get_id':
                        func.argtypes = [ctypes.c_short]
                        func.restype = ctypes.c_short
                    elif key == 'get_byte':
                        func.argtypes = [ctypes.c_short]
                        func.restype = ctypes.c_ubyte
                    elif key == 'set_dev_type_srch':
                        func.argtypes = [ctypes.c_short]
                        func.restype = ctypes.c_ushort
                    elif key == 'set_connect_product':
                        func.argtypes = [ctypes.c_uint]
                        func.restype = ctypes.c_ushort
                    elif key == 'get_product':
                        func.argtypes = []
                        func.restype = ctypes.c_uint
                    
                    exposed[key] = func
                    print(f"  Found: {func_name}")
                    break
                except AttributeError:
                    pass
        
        # Try to load defaults (optional)
        for func_name in ['pcProxPlusDefaults']:
            try:
                func = getattr(self.dll, func_name)
                print(f"  Found: {func_name}")
                # Call it to set defaults
                try:
                    func()
                except:
                    pass
                break
            except AttributeError:
                pass
        
        self.com_connect = exposed.get('com_connect')
        self.com_connect_port = exposed.get('com_connect_port')
        self.usb_disconnect = exposed.get('usb_disconnect')
        self.com_disconnect = exposed.get('com_disconnect')
        self.ping = exposed.get('ping')
        self.get_dev_cnt = exposed.get('get_dev_cnt')
        self.set_act_dev = exposed.get('set_act_dev')
        self.read_cfg = exposed.get('read_cfg')
        self.set_dev_type_srch = exposed.get('set_dev_type_srch')
        self.set_connect_product = exposed.get('set_connect_product')
        self.get_product = exposed.get('get_product')
        self.get_id_func = exposed.get('get_id')
        self.get_byte_func = exposed.get('get_byte')
        
        if not self.get_id_func or not self.get_byte_func:
            print(f"\n✗ Missing critical functions:")
            print(f"  get_id_func: {self.get_id_func}")
            print(f"  get_byte_func: {self.get_byte_func}")
            raise Exception("Could not find getActiveID32 or getActiveID_byte functions in pcProxAPI.dll")
        
        # Try to initialize connection
        self._initialize_device()
        
        print("✓ RF IDeas SDK initialized")
    
    def _initialize_device(self):
        """Initialize SDK connection to device"""
        try:
            # First disconnect any existing connections
            print("  Resetting SDK connections...")
            if self.usb_disconnect:
                try:
                    result = self.usb_disconnect()
                    print(f"  usbDisconnect result: {result}")
                except:
                    pass
            
            if self.com_disconnect:
                try:
                    result = self.com_disconnect()
                    print(f"  comDisconnect result: {result}")
                except:
                    pass
            
            time.sleep(0.5)
            
            # Try to search for device by type
            if self.set_dev_type_srch:
                # -1 = All types, 0 = USB, 1 = Serial, 2 = TCP/IP
                for dev_type in [-1, 0, 2]:  # Try All, USB, TCP/IP
                    result = self.set_dev_type_srch(dev_type)
                    print(f"  setDevTypeSrch({dev_type}) result: {result}")
                    time.sleep(0.1)
            
            # Connect to COM device
            if self.com_connect:
                result = self.com_connect()
                print(f"  comConnect result: {result}")
            
            # Test if device responds to ping
            if self.ping:
                try:
                    result = self.ping()
                    print(f"  Ping result: {result} (device is alive)")
                except Exception as e:
                    print(f"  Ping error: {e}")
            
            # Read configuration
            if self.read_cfg:
                result = self.read_cfg()
                print(f"  readCfg result: {result}")
            
            # Get number of devices
            if self.get_dev_cnt:
                dev_cnt = self.get_dev_cnt()
                print(f"  Device count: {dev_cnt}")
                
                if dev_cnt > 0 and self.set_act_dev:
                    result = self.set_act_dev(0)  # Select first device
                    print(f"  setActDev(0) result: {result}")
                elif dev_cnt == 0:
                    print(f"  No devices found via enumeration. Trying alternative access...")
                    # Try setting active device to 0 anyway - device may be accessed directly
                    if self.set_act_dev:
                        try:
                            result = self.set_act_dev(0)
                            print(f"  setActDev(0) result (forced): {result}")
                        except:
                            pass
                    
                    if self.set_connect_product:
                        # PRODUCT_ALL = 0x0FFFFFFFF
                        result = self.set_connect_product(0x0FFFFFFFF)
                        print(f"  setConnectProduct(ALL) result: {result}")
                        # Try again to get device count
                        dev_cnt = self.get_dev_cnt()
                        print(f"  Device count after setConnectProduct: {dev_cnt}")
        except Exception as e:
            print(f"  Warning during device initialization: {e}")

    def get_card(self):
        """Attempt to read a card once. Returns string UID or None"""
        if not self.get_id_func or not self.get_byte_func:
            return None
        
        try:
            # Call getActiveID32(256) to get data into internal buffer
            # Returns length of data or 0 if no card
            result = self.get_id_func(self.buffer_size)
            
            if result > 0:
                # Read individual bytes from the internal buffer
                uid = ""
                for i in range(result):
                    byte_val = self.get_byte_func(i)
                    uid += f"{byte_val:02X}"
                
                if uid:
                    return uid
        except Exception as e:
            # Silently ignore read errors - device may not have data
            pass
        return None

    def start_polling(self, output_queue, stop_event):
        """Continuously poll for cards and push into queue"""
        self.running = True
        while not stop_event.is_set():
            card = self.get_card()
            if card:
                output_queue.put(card)
                time.sleep(0.5)  # Prevent duplicate spam
            time.sleep(0.05)
        self.running = False

    def close(self):
        try:
            if self.end_func:
                self.end_func()
        except:
            pass

# --- CONFIGURATION ---
NUM_CARDS = 16777216
START_VALUE = 0x000705
STEP_SIZE   = 0x000001
KEY_TYPE = "EM4100/32"
UID_LENGTH = 10
EMULATION_DELAY = 0.0
DELAY_BETWEEN_CARDS = 0.0
DETECTION_TIMEOUT = 1.0
READER_RETRIES = 20
RESET_INTERVAL = 40
EXCEL_FILE = "emulator_log_w_sdk.xlsx"
LOCK_FILE = EXCEL_FILE + ".lock"

print("Starting emulator...")

# --- PORT DETECTION FUNCTIONS ---
def find_port_by_vid_pid(vid, pid):
    """Find serial port by VID and PID"""
    ports = list_ports.comports()
    for port in ports:
        if port.vid is not None and port.pid is not None:
            if port.vid == vid and port.pid == pid:
                return port.device
    return None

def check_flipper_device():
    """Check and initialize Flipper device"""
    flipper_port = find_port_by_vid_pid(0x0483, 0x5740)  # STMicroelectronics Flipper Zero
    if not flipper_port:
        print("✗ Flipper not found. Available ports:")
        for port in list_ports.comports():
            print(f"  {port.device}: {port.description} (VID: {port.vid:04X}, PID: {port.pid:04X})")
        flipper_port = input("Enter Flipper port manually: ") or 'COM17'
    else:
        print(f"✓ Flipper detected on {flipper_port}")
    
    try:
        flipper = serial.Serial(flipper_port, 230400, timeout=1)
        flipper.setDTR(True)
        flipper.setRTS(True)
        time.sleep(2)
        flipper.write(b"\r\n\r\n")  # wake up Flipper CLI
        time.sleep(0.5)
        print(f"✓ Flipper initialized on {flipper_port}")
        return flipper
    except Exception as e:
        print(f"✗ Error initializing Flipper: {e}")
        exit(1)

def check_reader_device():
    """Check and initialize helloID Reader device"""
    reader_port = find_port_by_vid_pid(0x1d6b, 0x4944)  # helloID Reader device
    if not reader_port:
        print("✗ Reader not found. Available ports:")
        for port in list_ports.comports():
            print(f"  {port.device}: {port.description} (VID: {port.vid:04X}, PID: {port.pid:04X})")
        reader_port = input("Enter Reader port manually: ") or 'COM18'
    else:
        print(f"✓ Reader detected on {reader_port}")
    
    try:
        reader = serial.Serial(reader_port, 9600, timeout=0)
        print(f"✓ Reader initialized on {reader_port}")
        return reader
    except Exception as e:
        print(f"✗ Error initializing Reader: {e}")
        exit(1)

def check_rfideas_device():
    """Check and initialize RFIDEAs device via pcProxAPI.dll SDK"""
    try:
        rfideas = RFIdeasReader()
        return rfideas
    except Exception as e:
        print(f"✗ Error initializing RF IDeas reader: {e}")
        return None

# --- SERIAL SETUP ---
flipper = check_flipper_device()
reader = check_reader_device()
rfideas = check_rfideas_device()

# --- QUEUES ---
reader_queue = queue.Queue()
keyboard_queue = queue.Queue()
stop_event = threading.Event()  

# --- HELPER FUNCTIONS ---
def generate_24bit_numbers(start=START_VALUE, step=STEP_SIZE):
    """Generate 24-bit numbers sequentially with the given step size"""
    if step <= 0:
        raise ValueError("Step must be a positive integer")
    max_value = 0xFFFFFF
    for value in range(start, max_value + 1, step):
        yield value

def convert24(value: int) -> int:
    """Convert 24-bit value using the convert24 algorithm"""
    value &= 0xFFFFFF
    high = (value >> 16) & 0xFF
    folded_high = (high & 0x80) | (high & 0x07)
    out = (value & 0x00FFFF) | (folded_high << 16)
    
    if (folded_high & 0x07) >= 4:
        out = (out - 0x01042E) & 0xFFFFFF

    return out

def int24_to_hex6(value: int) -> str:
    """Convert 24-bit integer to 6-character uppercase hex string"""
    value &= 0xFFFFFF
    
    return f"{value:06X}"
    
def clean_reader_line(line):
    """Extract UID from reader line. Keep only the last part after ':'"""
    if ':' in line:
        return line.split(':')[-1].strip()
    return line.strip()

def log_to_excel(nr, emulated, helloid, converted, rfideas, compare_result):
    """Log a single row to Excel with all columns."""
    with FileLock(LOCK_FILE):
        try:
            wb = load_workbook(EXCEL_FILE)
            ws = wb.active
        except FileNotFoundError:
            wb = Workbook()
            ws = wb.active
            ws.append(["Nr", "Emulated UID", "HelloID", "Converted HelloID", "RFIDEAs", "Compare"])
        ws.append([nr, emulated, helloid, converted, rfideas, compare_result])
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(EXCEL_FILE)

def send_flipper_command(key_data):
    key_type = KEY_TYPE
    cmd = f"rfid emulate {key_type} {key_data}"
    flipper.write((cmd + "\r\n").encode())
    time.sleep(1)

def stop_flipper_emulation():
    """Send Ctrl+C to stop emulation"""
    flipper.write(b"\x03")  # ASCII: Ctrl+C
    time.sleep(0.5)

def reset_flipper():
    """Reset the Flipper device by closing and reopening the connection"""
    global flipper
    try:
        flipper.close()
        time.sleep(1)
        # Re-detect Flipper port
        flipper_port = find_port_by_vid_pid(0x0483, 0x5740)  # STMicroelectronics Flipper Zero
        if not flipper_port:
            print("✗ Flipper not found after reset attempt")
            return
        flipper = serial.Serial(flipper_port, 230400, timeout=1)
        flipper.setDTR(True)
        flipper.setRTS(True)
        time.sleep(2)
        flipper.write(b"\r\n\r\n")  # Wake up Flipper CLI
        time.sleep(0.5)
        reader.reset_input_buffer()  # Clear reader input buffer during reset
        print(f"✓ Flipper reset")
    except Exception as e:
        print(f"✗ Error resetting Flipper: {e}")

# --- READER THREAD ---
def reader_thread():
    buffer = ""
    try:
        while True:
            bytes_waiting = reader.in_waiting
            if bytes_waiting > 0:
                data = reader.read(bytes_waiting).decode(errors='ignore')
                buffer += data
                if "\n" in buffer:
                    lines = buffer.split("\n")
                    for line in lines[:-1]:
                        line = line.strip()
                        if line:
                            uid = clean_reader_line(line)
                            
                            # Validate before queuing
                            if len(uid) >= 6:  # Minimum UID length
                                try:
                                    int(uid, 16)  # Check if valid hex
                                    reader_queue.put(uid)
                                except ValueError:
                                    pass  # Skip invalid hex
                            # else skip short/truncated reads like "000"
                    buffer = lines[-1]
            time.sleep(0.05)
    except KeyboardInterrupt:
        print("\nExiting reader thread...")
        reader.close()

# --- RF IDeas POLLING THREAD ---
def rfideas_polling_thread():
    """Poll RF IDeas SDK for card reads"""
    if rfideas:
        rfideas.start_polling(keyboard_queue, stop_event)

# --- GENERATE EMULATED CARDS ---
emulated_cards = []
for num in generate_24bit_numbers(step=STEP_SIZE):
    euid = f"00000012D6{num:06X}"
    emulated_cards.append(euid)
    if len(emulated_cards) >= NUM_CARDS:
        break

# --- MAIN LOOP ---
def main():
    # Determine starting Nr from Excel
    try:
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        last_row = ws.max_row
        if last_row > 1:  # Header is row 1
            last_nr = ws.cell(row=last_row, column=1).value
            start_nr = last_nr + 1 if isinstance(last_nr, int) else 1
        else:
            start_nr = 1
    except FileNotFoundError:
        start_nr = 1
    except Exception:
        start_nr = 1

    # Print headers
    print("\n")
    print(f"{'Nr':<5}{'Emulated UID':<20}{'HelloID':<20}{'Converted HelloID':<20}{'RFIDEAs':<20}{'Compare'}")
    print("-"*110)

    for i, euid in enumerate(emulated_cards, start_nr):
        # Reset Flipper at intervals to prevent device timeout
        if (i - start_nr) > 0 and (i - start_nr) % RESET_INTERVAL == 0:
            reset_flipper()
            time.sleep(1)
        
        # Clear old queue items BEFORE emulation
        while not reader_queue.empty():
            reader_queue.get()
        while not keyboard_queue.empty():
            keyboard_queue.get()
        time.sleep(0.2)  # Wait for queues to fully clear

        send_flipper_command(euid[-10:])
        time.sleep(EMULATION_DELAY)  # Give devices time to respond after emulation

        # Wait for helloID reader input with retry
        reader_read = "NOT DETECTED"
        for attempt in range(READER_RETRIES):
            try:
                reader_read = reader_queue.get(timeout=DETECTION_TIMEOUT)
                break
            except queue.Empty:
                if attempt < READER_RETRIES - 1:  # Don't sleep after last attempt
                    time.sleep(0.5)  # Short delay before retry

        # Wait for keyboard/RFIDEAs input
        rfideas_read = ""
        try:
            rfideas_read = keyboard_queue.get(timeout=DETECTION_TIMEOUT)
        except queue.Empty:
            rfideas_read = "NOT DETECTED"

        time.sleep(0.2)
        stop_flipper_emulation()

        # Convert the received UID from reader (last 6 hex chars = 3 bytes)
        converted_uid = "NOT DETECTED"
        if reader_read != "NOT DETECTED":
            try:
                # Extract last 6 characters, convert them, keep first 10 chars of device ID
                int_value = int(reader_read[-6:], 16)
                converted_int = convert24(int_value)
                hex_value = int24_to_hex6(converted_int)
                # reader_read is like "00000012D6000072" (16 chars: 10 prefix + 6 uid)
                # Replace only the last 6 chars with converted value
                converted_uid = reader_read[:10] + hex_value
            except (ValueError, IndexError):
                converted_uid = "ERROR"

        # Format all UIDs with 000000 prefix for display
        emulated_display = euid
        helloid_display = reader_read if reader_read != "NOT DETECTED" else "NOT DETECTED"
        converted_display = converted_uid if converted_uid not in ["NOT DETECTED", "ERROR"] else converted_uid
        rfideas_display = "000000" + rfideas_read if rfideas_read != "NOT DETECTED" else "NOT DETECTED"
        
        # Compare Converted HelloID with RFIDEAs
        if converted_display in ["NOT DETECTED", "ERROR"] or rfideas_display == "NOT DETECTED":
            compare_result = "ERROR"
        elif converted_display == rfideas_display:
            compare_result = "Matched"
        else:
            compare_result = "Not Matched"

        # Log and print: Nr | Emulated UID | HelloID | Converted HelloID | RFIDEAs | Compare
        print(f"{i:<5}{emulated_display:<20}{helloid_display:<20}{converted_display:<20}{rfideas_display:<20}{compare_result}")
        log_to_excel(i, emulated_display, helloid_display, converted_display, rfideas_display, compare_result)

        time.sleep(DELAY_BETWEEN_CARDS)

# --- MAIN ---
if __name__ == "__main__":
    threading.Thread(target=reader_thread, daemon=True).start()
    threading.Thread(target=rfideas_polling_thread, daemon=True).start()
    time.sleep(1)
    main()
    print("\nAll cards emulated and logged to:", EXCEL_FILE)
   
    # Signal stop and cleanup
    stop_event.set()
    if rfideas:
        rfideas.close()
    time.sleep(0.1)
    exit(0)
    