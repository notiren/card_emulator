import serial
import time
from openpyxl import load_workbook, Workbook
from filelock import FileLock
from datetime import datetime

# --- CONFIG ---
READER_PORT = input("Enter Reader port (e.g., COM18): ") or 'COM18'
READER_BAUD = 9600
EXCEL_FILE = "combined_log.xlsx"
LOCK_FILE = EXCEL_FILE + ".lock"

# --- SETUP SERIAL ---
reader = serial.Serial(READER_PORT, READER_BAUD, timeout=0)
print(f"Listening on {READER_PORT} at {READER_BAUD} baud...")

def log_to_excel(source, data):
    """Safely log to shared Excel file."""
    with FileLock(LOCK_FILE):
        try:
            wb = load_workbook(EXCEL_FILE)
            ws = wb.active
        except FileNotFoundError:
            wb = Workbook()
            ws = wb.active
            ws.append(["Timestamp", "Source", "Data"])
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        ws.append([timestamp, source, data])
        wb.save(EXCEL_FILE)

# --- READ LOOP ---
buffer = ""
try:
    while True:
        bytes_waiting = reader.in_waiting
        if bytes_waiting > 0:
            data = reader.read(bytes_waiting).decode(errors='ignore')
            buffer += data
            if "\n" in buffer:
                lines = buffer.split("\n")
                for line in lines[:-1]:
                    line = line.strip()
                    if line:
                        print(line)
                        log_to_excel("Reader", line)
                buffer = lines[-1]  # keep partial line
        time.sleep(0.05)
except KeyboardInterrupt:
    print("\nExiting reader...")
    reader.close()