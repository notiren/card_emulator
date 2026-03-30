import serial
import time
import random
from openpyxl import load_workbook, Workbook
from filelock import FileLock
from datetime import datetime

# --- CONFIG ---
FLIPPER_PORT = input("Enter Flipper port (e.g., COM17): ") or 'COM17'
FLIPPER_BAUD = 230400
NUM_CARDS = 50
DELAY_BETWEEN_CARDS = 5
EXCEL_FILE = "combined_log.xlsx"
LOCK_FILE = EXCEL_FILE + ".lock"

# --- SETUP FLIPPER SERIAL ---
flipper = serial.Serial(FLIPPER_PORT, FLIPPER_BAUD, timeout=1)
flipper.setDTR(True)
flipper.setRTS(True)
time.sleep(6)  # allow CLI to start
flipper.write(b"\r\n\r\n")
time.sleep(0.5)

# --- HELPER FUNCTIONS ---
def generate_em4100_id():
    """Generate a random 10-digit hexadecimal EM4100 card ID."""
    return ''.join(random.choices("0123456789ABCDEF", k=10))

def send_flipper_command(cmd):
    """Send command to Flipper CLI (ignore response)."""
    flipper.write((cmd + "\r\n").encode())
    time.sleep(1)

def log_to_excel(source, data):
    """Safely log to shared Excel file."""
    with FileLock(LOCK_FILE):
        try:
            wb = load_workbook(EXCEL_FILE)
            ws = wb.active
        except FileNotFoundError:
            wb = Workbook()
            ws = wb.active
            ws.append(["Timestamp", "Source", "Data"])
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        ws.append([timestamp, source, data])
        wb.save(EXCEL_FILE)

# --- MAIN LOOP ---
for i in range(NUM_CARDS):
    key_type = "EM4100/32"
    key_data = generate_em4100_id()
    cmd = f"rfid emulate {key_type} {key_data}"

    print(f"[{i+1}/{NUM_CARDS}] Emulating card: {key_data}")
    send_flipper_command(cmd)
    log_to_excel("Flipper", key_data)  # log only the sent card

    time.sleep(DELAY_BETWEEN_CARDS)

print("All cards emulated. Log saved to:", EXCEL_FILE)